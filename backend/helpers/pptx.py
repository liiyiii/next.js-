# coding=utf-8
from pptx.oxml.xmlchemy import OxmlElement
from pptx.dml.color import RGBColor
from pptx.util import Cm, Mm, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

"""lib"""


def sub_element(parent, tagname, **kwargs):
    """create tag element with given attributes,
    then add it to given parent."""
    element = OxmlElement(tagname)
    element.attrib.update(kwargs)
    parent.append(element)
    return element


"""set style"""


def make_paragraph_bullet_pointed(para, margin_left=171450, indent=171450):
    """Bullets are set to Arial,
        actual text can be a different font"""
    pPr = para._p.get_or_add_pPr()
    # Set marL and indent attributes
    pPr.set('marL', str(margin_left))  # margin left
    pPr.set('indent', str(indent))  # indent
    # Add buFont - bulletin font
    _ = sub_element(parent=pPr, tagname="a:buFont", typeface="Arial")
    # Add buChar - dot
    _ = sub_element(parent=pPr, tagname='a:buChar', char="•")


def set_cell_border(cell, border_color="000000", border_width='12700'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()  # table cell property
    for lines in ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']:
        # get boder line
        ln = sub_element(tcPr, lines, w=border_width, cap='flat', cmpd='sng', algn='ctr')
        # set line color
        solidFill = sub_element(ln, 'a:solidFill')
        srgbClr = sub_element(solidFill, 'a:srgbClr', val=border_color)  # boder color
        prstDash = sub_element(ln, 'a:prstDash', val='solid')  # dash style
        round_ = sub_element(ln, 'a:round')  # round
        headEnd = sub_element(ln, 'a:headEnd', type='none', w='med', len='med')  # arrow left
        tailEnd = sub_element(ln, 'a:tailEnd', type='none', w='med', len='med')  # arrow right


"""dump content"""


def write_image(shapes, image_path, xy=(0, 0), width=0, height=0):
    # Ensure coordinates are converted to Pt if they are simple numbers
    x_pt = xy[0] if isinstance(xy[0], Pt) else Pt(xy[0])
    y_pt = xy[1] if isinstance(xy[1], Pt) else Pt(xy[1])

    if width != 0:
        width_pt = width if isinstance(width, Pt) else Pt(width)
        shapes.add_picture(image_path, x_pt, y_pt, width=width_pt) # x, y are often 0,0 for full slide image
    elif height != 0:
        height_pt = height if isinstance(height, Pt) else Pt(height)
        shapes.add_picture(image_path, x_pt, y_pt, height=height_pt) # x, y are often 0,0 for full slide image
    else:
        shapes.add_picture(image_path, x_pt, y_pt)


def write_text(shapes, text, xy, size, font_size=9, font_name='Arial'):
    textbox = shapes.add_textbox(Pt(xy[0]), Pt(xy[1]), Pt(size[0]), Pt(size[1]))
    tf = textbox.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(font_size)


def write_text_block(shapes, text, xy, size, name='AITextBox', word_warp=True,
                     font_size=9, font_name='Arial',
                     font_color=None, bgcolor=None, framecolor=None,
                     bold=False, alignment='left'):
    # from pptx.enum.text import MSO_AUTO_SIZE # Already imported at top

    textbox = shapes.add_textbox(Pt(xy[0]), Pt(xy[1]), Pt(size[0]), Pt(size[1]))
    textbox.name = name

    if framecolor is not None:
        try:
            textbox.line.color.rgb = RGBColor(r=framecolor[0], g=framecolor[1], b=framecolor[2])
        except Exception as e:
            print(f"Error setting framecolor: {e}. Framecolor was: {framecolor}")


    if bgcolor is not None:
        try:
            textbox.fill.solid()
            textbox.fill.fore_color.rgb = RGBColor(r=bgcolor[0], g=bgcolor[1], b=bgcolor[2])
        except Exception as e:
            print(f"Error setting bgcolor: {e}. Bgcolor was: {bgcolor}")

    tf = textbox.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    if word_warp:
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    p = tf.paragraphs[0]
    p.text = text
    p.font.bold = bold
    p.font.name = font_name
    p.font.size = Pt(font_size)
    if font_color is not None:
        try:
            color = RGBColor(r=font_color[0], g=font_color[1], b=font_color[2])
            p.font.color.rgb = color
        except Exception as e:
            print(f"Error setting font_color: {e}. Font_color was: {font_color}")


    if alignment == 'center':
        p.alignment = PP_ALIGN.CENTER
    elif alignment == 'left':
        p.alignment = PP_ALIGN.LEFT
    elif alignment == 'right':
        p.alignment = PP_ALIGN.RIGHT


def write_list_block(shapes, text_list, xy, size, name='AITextBox', font_size=9, font_name='Arial',
                     font_color=None, bgcolor=None, framecolor=None):
    textbox = shapes.add_textbox(Pt(xy[0]), Pt(xy[1]), Pt(size[0]), Pt(size[1]))
    textbox.name = name

    if framecolor is not None:
        textbox.line.color.rgb = RGBColor(r=framecolor[0], g=framecolor[1], b=framecolor[2])

    if bgcolor is not None:
        textbox.fill.solid()
        textbox.fill.fore_color.rgb = RGBColor(r=bgcolor[0], g=bgcolor[1], b=bgcolor[2])

    tf = textbox.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    for i, text_item in enumerate(text_list):
        if i == 0:
            p = tf.paragraphs[0]  # first paragraph always there
        else:
            p = tf.add_paragraph()
        make_paragraph_bullet_pointed(p, indent=Pt(font_size).emu)  # use bullet

        r = p.add_run()
        r.text = text_item # Corrected from text to text_item
        r.font.name = font_name
        r.font.size = Pt(font_size)  # define size with point metric
        if font_color: # Check if font_color is provided
            color = RGBColor(r=font_color[0], g=font_color[1], b=font_color[2])
            r.font.color.rgb = color  # define color


def write_simple_table(shapes, bbox, contents, font_size=9, font_color=None, bg_color=None, border_color="CCCCCC"):
    nrow = len(contents)
    if nrow == 0: return # No rows to write
    ncol = len(contents[0])
    if ncol == 0: return # No columns to write


    # create table
    table_shape = shapes.add_table(
        nrow, ncol, Pt(bbox[0]), Pt(bbox[1]),
        Pt(bbox[2]-bbox[0]), Pt(bbox[3]-bbox[1]))

    table = table_shape.table # Get the Table object
    table.first_row = False # No special formatting for first row by default

    # rendering
    for i in range(nrow):
        for j in range(ncol):
            cell = table.cell(i, j)
            set_cell_border(cell, border_color=border_color)

            # Ensure paragraph exists, python-pptx usually creates one
            if not cell.text_frame.paragraphs:
                 p = cell.text_frame.add_paragraph()
            else:
                 p = cell.text_frame.paragraphs[0]

            p.text = str(contents[i][j]) # Ensure content is string
            p.font.name = 'Arial'
            p.font.size = Pt(font_size)

            if font_color is not None:
                # Assuming font_color is (R, G, B) tuple
                p.font.color.rgb = RGBColor(font_color[0], font_color[1], font_color[2])

            # p.alignment = PP_ALIGN.CENTER # Default or make it a parameter

            # bgcolor
            if bg_color is not None:
                cell.fill.solid()
                # Assuming bg_color is (R, G, B) tuple
                cell.fill.fore_color.rgb = RGBColor(bg_color[0], bg_color[1], bg_color[2])
